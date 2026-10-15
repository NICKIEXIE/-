#!/usr/bin/env python3
"""按 JSON 生成 Word 交付物（起号方案书 / 脚本表 / 验收单 / 复盘报告）。

JSON 结构（通用）：

{
  "title": "必填",
  "subtitle": "可选",
  "meta": [{"label": "客户", "value": "XX 餐饮"}],          // 可选，渲染成信息表
  "sections": [
    {
      "heading": "必填",
      "level": 1,                                            // 1=Heading1, 2=Heading2, 默认 1
      "paragraphs": ["正文段落"],
      "bullets": ["要点"],
      "table": {"headers": ["列1","列2"], "rows": [["a","b"]]},
      "note": "提示框文字"
    }
  ],
  "footer": "可选"
}

用法：
  python make_docx.py --kind plan --input plan.json --output 起号方案书.docx
  python make_docx.py --kind script --input script.json --output 脚本表.docx
  python make_docx.py --kind acceptance --input acc.json --output 验收单.docx
  python make_docx.py --kind report --input report.json --output 复盘报告.docx
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:  # pragma: no cover
    pass

BODY_FONT = "微软雅黑"
MONO_FONT = "等线"

KINDS: dict[str, dict] = {
    "plan": {
        "default_title": "短视频起号方案书",
        "min_sections": 4,
        "require_meta": 3,
        "hint": "方案书至少 4 个章节，且 meta 至少 3 项（客户/行业/目标等）",
    },
    "script": {
        "default_title": "视频脚本表",
        "min_sections": 1,
        "require_meta": 0,
        "hint": "每个章节代表一条视频，章节内用 table 或 bullets 给出钩子/正文/结尾",
    },
    "acceptance": {
        "default_title": "代运营交付验收单",
        "min_sections": 2,
        "require_meta": 2,
        "hint": "验收单至少 2 个章节，并含一张至少 2 列的验收项表格",
    },
    "report": {
        "default_title": "运营复盘报告",
        "min_sections": 2,
        "require_meta": 1,
        "hint": "报告至少 2 个章节（如现状/诊断/动作）",
    },
}


class SpecError(Exception):
    """JSON 结构不符合约定。"""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise SpecError(message)


def load_spec(path: Path) -> dict:
    try:
        raw = path.read_text(encoding="utf-8")
    except OSError as error:
        raise SpecError(f"读不到输入文件 {path}: {error}") from error
    try:
        spec = json.loads(raw)
    except json.JSONDecodeError as error:
        raise SpecError(f"输入不是合法 JSON（第 {error.lineno} 行第 {error.colno} 列）: {error.msg}") from error
    require(isinstance(spec, dict), "JSON 顶层必须是对象 {}")
    return spec


def validate(spec: dict, kind: str) -> None:
    rules = KINDS[kind]
    title = spec.get("title")
    require(isinstance(title, str) and title.strip(), "缺少 title（非空字符串）")

    meta = spec.get("meta", [])
    require(isinstance(meta, list), "meta 必须是数组")
    if rules["require_meta"]:
        require(
            len(meta) >= rules["require_meta"],
            f"kind={kind} 的 meta 至少需要 {rules['require_meta']} 项（每项 {{\"label\":\"…\",\"value\":\"…\"}}）",
        )
    for index, item in enumerate(meta):
        require(isinstance(item, dict), f"meta[{index}] 必须是对象")
        require(isinstance(item.get("label"), str), f"meta[{index}].label 必须是字符串")
        require(isinstance(item.get("value"), str), f"meta[{index}].value 必须是字符串")

    sections = spec.get("sections")
    require(isinstance(sections, list) and sections, "sections 必须是非空数组")
    require(
        len(sections) >= rules["min_sections"],
        f"kind={kind} 至少需要 {rules['min_sections']} 个章节；{rules['hint']}",
    )

    has_two_column_table = False
    for index, section in enumerate(sections):
        require(isinstance(section, dict), f"sections[{index}] 必须是对象")
        require(
            isinstance(section.get("heading"), str) and section["heading"].strip(),
            f"sections[{index}].heading 必填（非空字符串）",
        )
        level = section.get("level", 1)
        require(level in (1, 2, 3), f"sections[{index}].level 只能是 1/2/3")
        body_keys = [key for key in ("paragraphs", "bullets", "table") if section.get(key)]
        require(body_keys, f"sections[{index}] 内容为空：paragraphs / bullets / table 至少要有一个")
        for key in ("paragraphs", "bullets"):
            if key in section and section[key] is not None:
                require(isinstance(section[key], list), f"sections[{index}].{key} 必须是数组")
                require(
                    all(isinstance(item, str) for item in section[key]),
                    f"sections[{index}].{key} 的每一项都必须是字符串",
                )
        if "table" in section and section["table"] is not None:
            table = section["table"]
            require(isinstance(table, dict), f"sections[{index}].table 必须是对象")
            headers = table.get("headers")
            rows = table.get("rows", [])
            require(isinstance(headers, list) and headers, f"sections[{index}].table.headers 必须是非空数组")
            require(isinstance(rows, list), f"sections[{index}].table.rows 必须是数组")
            for row_index, row in enumerate(rows):
                require(isinstance(row, list), f"sections[{index}].table.rows[{row_index}] 必须是数组")
                require(
                    len(row) == len(headers),
                    f"sections[{index}].table.rows[{row_index}] 有 {len(row)} 列，表头有 {len(headers)} 列，不一致",
                )
            if len(headers) >= 2:
                has_two_column_table = True

    if kind == "acceptance":
        require(has_two_column_table, "验收单至少需要一张 2 列以上的验收项表格")


def set_cjk_font(document: Document) -> None:
    """同时设置 ascii 与 eastAsia 字体，避免中文在 Word 里回落成别的字体。"""
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.font.name = BODY_FONT
            rpr = run._element.get_or_add_rPr()
            rfonts = rpr.get_or_add_rFonts()
            rfonts.set(qn("w:eastAsia"), BODY_FONT)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = BODY_FONT
                        rpr = run._element.get_or_add_rPr()
                        rfonts = rpr.get_or_add_rFonts()
                        rfonts.set(qn("w:eastAsia"), BODY_FONT)


def render(spec: dict, kind: str) -> Document:
    rules = KINDS[kind]
    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    try:
        rpr = normal.element.get_or_add_rPr()
        rpr.get_or_add_rFonts().set(qn("w:eastAsia"), BODY_FONT)
    except AttributeError:  # pragma: no cover - 样式元素结构变化时退化为逐 run 设置
        pass

    title = document.add_heading(spec.get("title") or rules["default_title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if spec.get("subtitle"):
        subtitle = document.add_paragraph(str(spec["subtitle"]))
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = spec.get("meta") or []
    if meta:
        table = document.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        for item in meta:
            cells = table.add_row().cells
            cells[0].text = item["label"]
            cells[1].text = item["value"]
        document.add_paragraph()

    for section in spec["sections"]:
        document.add_heading(section["heading"], level=int(section.get("level", 1)))
        for paragraph in section.get("paragraphs") or []:
            document.add_paragraph(paragraph)
        for bullet in section.get("bullets") or []:
            document.add_paragraph(bullet, style="List Bullet")
        table_spec = section.get("table")
        if table_spec:
            headers = table_spec["headers"]
            table = document.add_table(rows=1, cols=len(headers))
            table.style = "Light Grid Accent 1"
            for index, header in enumerate(headers):
                cell = table.rows[0].cells[index]
                cell.text = str(header)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
            for row in table_spec.get("rows", []):
                cells = table.add_row().cells
                for index, value in enumerate(row):
                    cells[index].text = "" if value is None else str(value)
            document.add_paragraph()
        if section.get("note"):
            note = document.add_paragraph()
            run = note.add_run(f"提示：{section['note']}")
            run.italic = True
            run.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    if spec.get("footer"):
        document.add_paragraph()
        footer = document.add_paragraph(str(spec["footer"]))
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    set_cjk_font(document)
    return document


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="按 JSON 生成 Word 交付物")
    parser.add_argument("--kind", required=True, choices=sorted(KINDS), help="交付物类型")
    parser.add_argument("--input", required=True, help="输入 JSON 路径")
    parser.add_argument("--output", required=True, help="输出 .docx 路径（必须是新文件）")
    parser.add_argument("--force", action="store_true", help="允许覆盖已存在的输出文件")
    args = parser.parse_args(argv[1:])

    source = Path(args.input).expanduser().resolve()
    target = Path(args.output).expanduser().resolve()
    if target.suffix.lower() != ".docx":
        print(f"输出必须是 .docx：{target}")
        return 2
    if target.exists() and not args.force:
        print(f"输出文件已存在，请换名字或用 --force：{target}")
        return 2

    try:
        spec = load_spec(source)
        validate(spec, args.kind)
    except SpecError as error:
        print(f"JSON 不符合约定：{error}")
        return 2

    document = render(spec, args.kind)
    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))

    paragraphs = len(document.paragraphs)
    print(f"已生成：{target}")
    print(f"kind={args.kind}  章节={len(spec['sections'])}  段落={paragraphs}  表格={len(document.tables)}  大小={target.stat().st_size} 字节")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
