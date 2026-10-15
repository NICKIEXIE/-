#!/usr/bin/env python3
"""按 JSON 生成 Excel 交付物（排期表 / 报价单 / 需求诊断表 / 数据跟踪表）。

JSON 结构：

{
  "title": "可选，仅用于日志",
  "sheets": [
    {
      "name": "30天排期",                      // 必填，工作表名（≤31 字符，不能含 []:*?/\\）
      "title": "起号期 30 天排期表",             // 可选，写在第 1 行并合并
      "headers": ["第几天", "日期", "内容主题", "动作", "负责人"],   // 必填
      "rows": [["D1", "10-20", "自我介绍", "拍摄+剪辑", "老板"]],
      "widths": [8, 12, 24, 24, 10],           // 可选，不填按内容自动估
      "notes": ["发布后 30 分钟内回复全部评论"]   // 可选，写在表下方
    }
  ]
}

用法：
  python make_xlsx.py --kind schedule --input schedule.json --output 30天排期表.xlsx
  python make_xlsx.py --kind quote    --input quote.json    --output 报价单.xlsx
  python make_xlsx.py --kind diagnosis --input diag.json    --output 需求诊断表.xlsx
  python make_xlsx.py --kind metrics  --input metrics.json  --output 数据跟踪表.xlsx
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:  # pragma: no cover
    pass

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(name="微软雅黑", size=10.5, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10.5)
TITLE_FONT = Font(name="微软雅黑", size=13, bold=True)
NOTE_FONT = Font(name="微软雅黑", size=9, italic=True, color="808080")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")

KINDS: dict[str, dict] = {
    "schedule": {"hint": "排期表至少要有一列表头含“天”或“日期”", "header_keywords": ("天", "日期")},
    "quote": {"hint": "报价单至少要有一列表头含“报价”“单价”“金额”或“小计”", "header_keywords": ("报价", "单价", "金额", "小计")},
    "diagnosis": {"hint": "需求诊断表至少要有一列表头含“问题”或“需求”", "header_keywords": ("问题", "需求")},
    "metrics": {"hint": "数据跟踪表至少要有一列表头含“指标”或“数值”", "header_keywords": ("指标", "数值")},
}


class SpecError(Exception):
    """JSON 结构不符合约定。"""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise SpecError(message)


def load_spec(path: Path) -> dict:
    try:
        raw = path.read_text(encoding="utf-8")
    except OSError as error:
        raise SpecError(f"读不到输入文件 {path}: {error}") from error
    try:
        spec = json.loads(raw)
    except json.JSONDecodeError as error:
        raise SpecError(f"输入不是合法 JSON（第 {error.lineno} 行第 {error.colno} 列）: {error.msg}") from error
    require(isinstance(spec, dict), "JSON 顶层必须是对象 {}")
    return spec


def validate(spec: dict, kind: str) -> None:
    sheets = spec.get("sheets")
    require(isinstance(sheets, list) and sheets, "sheets 必须是非空数组")
    rules = KINDS[kind]

    matched = False
    for index, sheet in enumerate(sheets):
        prefix = f"sheets[{index}]"
        require(isinstance(sheet, dict), f"{prefix} 必须是对象")
        name = sheet.get("name")
        require(isinstance(name, str) and name.strip(), f"{prefix}.name 必填（工作表名）")
        require(len(name) <= 31, f"{prefix}.name 超过 31 字符（Excel 限制）")
        require(not INVALID_SHEET_CHARS.search(name), f"{prefix}.name 不能包含 [ ] : * ? / \\ 这些字符")

        headers = sheet.get("headers")
        require(isinstance(headers, list) and headers, f"{prefix}.headers 必须是非空数组")
        require(
            all(isinstance(item, str) and item.strip() for item in headers),
            f"{prefix}.headers 的每一项都必须是非空字符串",
        )
        rows = sheet.get("rows", [])
        require(isinstance(rows, list), f"{prefix}.rows 必须是数组")
        for row_index, row in enumerate(rows):
            require(isinstance(row, list), f"{prefix}.rows[{row_index}] 必须是数组")
            require(
                len(row) == len(headers),
                f"{prefix}.rows[{row_index}] 有 {len(row)} 列，表头 {len(headers)} 列，不一致",
            )
        widths = sheet.get("widths")
        if widths is not None:
            require(isinstance(widths, list), f"{prefix}.widths 必须是数组")
            require(len(widths) == len(headers), f"{prefix}.widths 有 {len(widths)} 项，表头 {len(headers)} 列，不一致")
            require(
                all(isinstance(item, (int, float)) and item > 0 for item in widths),
                f"{prefix}.widths 必须都是正数",
            )
        notes = sheet.get("notes")
        if notes is not None:
            require(isinstance(notes, list), f"{prefix}.notes 必须是数组")
            require(all(isinstance(item, str) for item in notes), f"{prefix}.notes 的每一项都必须是字符串")

        if any(keyword in header for header in headers for keyword in rules["header_keywords"]):
            matched = True

    require(matched, f"kind={kind}：{rules['hint']}")


def display_width(value: object) -> int:
    text = "" if value is None else str(value)
    width = 0
    for char in text:
        width += 2 if ord(char) > 0x2E80 else 1
    return width


def estimate_widths(headers: list[str], rows: list[list]) -> list[float]:
    widths = []
    for index, header in enumerate(headers):
        longest = display_width(header)
        for row in rows[:200]:
            longest = max(longest, display_width(row[index]))
        widths.append(min(max(longest + 3, 9), 60))
    return widths


def render(spec: dict, kind: str) -> tuple[Workbook, list[str]]:
    workbook = Workbook()
    workbook.remove(workbook.active)
    written: list[str] = []

    for sheet_spec in spec["sheets"]:
        sheet = workbook.create_sheet(title=sheet_spec["name"])
        written.append(sheet_spec["name"])
        headers = sheet_spec["headers"]
        rows = sheet_spec.get("rows", [])
        notes = sheet_spec.get("notes") or []
        row_cursor = 1

        if sheet_spec.get("title"):
            sheet.cell(row=1, column=1, value=str(sheet_spec["title"])).font = TITLE_FONT
            sheet.merge_cells(
                start_row=1, start_column=1, end_row=1, end_column=max(len(headers), 1)
            )
            row_cursor = 3

        header_row = row_cursor
        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=header_row, column=index, value=header)
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.row_dimensions[header_row].height = 22

        for offset, row in enumerate(rows, start=1):
            for index, value in enumerate(row, start=1):
                cell = sheet.cell(row=header_row + offset, column=index, value=value)
                cell.font = BODY_FONT
                cell.border = BORDER
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        note_row = header_row + len(rows) + 2
        for offset, note in enumerate(notes):
            cell = sheet.cell(row=note_row + offset, column=1, value=str(note))
            cell.font = NOTE_FONT

        widths = sheet_spec.get("widths") or estimate_widths(headers, rows)
        for index, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(index)].width = float(width)
        sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
        sheet.sheet_view.showGridLines = False

    return workbook, written


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description="按 JSON 生成 Excel 交付物")
    parser.add_argument("--kind", required=True, choices=sorted(KINDS), help="交付物类型")
    parser.add_argument("--input", required=True, help="输入 JSON 路径")
    parser.add_argument("--output", required=True, help="输出 .xlsx 路径（必须是新文件）")
    parser.add_argument("--force", action="store_true", help="允许覆盖已存在的输出文件")
    args = parser.parse_args(argv[1:])

    source = Path(args.input).expanduser().resolve()
    target = Path(args.output).expanduser().resolve()
    if target.suffix.lower() != ".xlsx":
        print(f"输出必须是 .xlsx：{target}")
        return 2
    if target.exists() and not args.force:
        print(f"输出文件已存在，请换名字或用 --force：{target}")
        return 2

    try:
        spec = load_spec(source)
        validate(spec, args.kind)
    except SpecError as error:
        print(f"JSON 不符合约定：{error}")
        return 2

    workbook, sheets = render(spec, args.kind)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(str(target))

    total_rows = sum(len(sheet.get("rows", [])) for sheet in spec["sheets"])
    print(f"已生成：{target}")
    print(f"kind={args.kind}  工作表={len(sheets)}（{'、'.join(sheets)}）  数据行={total_rows}  大小={target.stat().st_size} 字节")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
